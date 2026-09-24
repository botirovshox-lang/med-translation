"""Пробный фрагмент: документ больше подарка переводится одной страницей, а не отказом.

Пробная организация (регистрация при SIGNUP_FREE_PAGES) принесла файл больше
остатка страниц. Прежде — 402 и человек, ушедший ни с чем. Теперь объём
всего файла считается бесплатно, а в проект встаёт ОДИН случайный непустой
фрагмент подряд идущих строк на остаток, и списывается только он.
Сторожится главное:
  * фрагмент — подряд, непустой, не больше остатка; списан фрагмент,
    а объём файла назван отдельно (строка журнала `excerpt`);
  * организация, которой платили (пополнил администратор), фрагмента
    не получает — там прежний 402;
  * всё, что завело бы строки по ВСЕМУ файлу, на фрагменте закрыто:
    пересборка строк и разбор картинок (иначе остаток книги уехал бы даром);
  * «Перевести остальное» без пополнения — 402, после пополнения дописывает
    строки из хранимого исходника и снимает отметку фрагмента.
Ни одного вызова модели.
"""
import io
import os
import sys
import tempfile
from pathlib import Path

os.environ["APP_PASSWORD"] = "boot-password-1"
os.environ["AUTHORITY_CORPUS"] = "0"
os.environ["SIGNUP_ENABLED"] = "1"
for k in ("SMTP_HOST", "MAIL_FROM", "SMTP_USER"):
    os.environ.pop(k, None)
sys.path.insert(0, "backend")
import main
from starlette.testclient import TestClient

TMP = Path(tempfile.mkdtemp(prefix="mct-trial-"))
main.SOURCE_DIR = TMP / "sources"
main.SOURCE_DIR.mkdir(parents=True, exist_ok=True)
main.REIMPORT_DIR = TMP / "backups"
main.save_state = lambda *a, **k: None
main.STATE["users"], main.STATE["tenants"], main.STATE["audit"] = [], [], []
main.STATE["spend"], main.STATE["projects"] = {}, []
main._SESSIONS.clear(); main._LOGIN_FAILS.clear(); main._SIGNUP_FAILS.clear()
main.SIGNUP_FREE_PAGES, main.SIGNUP_TRIAL_USD = 1.0, 0.3
main.TENANT_MAX_PAGES = 0
fail = []


def check(cond, label):
    print(("  OK   " if cond else "  FAIL ") + label)
    if not cond:
        fail.append(label)


try:
    from docx import Document
except ImportError:
    print("python-docx не установлен — проверка пропущена")
    sys.exit(0)

c = TestClient(main.app)
H = lambda t: {"Authorization": "Bearer " + t}
MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
A = c.post("/api/auth/login", json={"login": "admin", "password": "boot-password-1"}).json()["token"]


def docx_of(n, words=30, tag="абзац"):
    d = Document()
    for i in range(n):
        d.add_paragraph(("%s %d " % (tag, i)) + " ".join("слово%d" % (j % 7) for j in range(words - 2)) + ".")
    b = io.BytesIO(); d.save(b)
    return b.getvalue()


def signup(email):
    r = c.post("/api/auth/register", json={"email": email, "password": "long-enough-1", "accept": True})
    u = main._user_by_email(email)
    code = main._issue_code(u, "verify")
    tok = c.post("/api/auth/verify", json={"email": email, "code": code}).json()["token"]
    return r.json()["tenant"], tok


def up(tok, raw, name="book.docx"):
    return c.post("/api/projects/upload", headers=H(tok), files={"file": (name, raw, MIME)},
                  data={"src": "RU", "tgt": "EN"})


print("=== 1. Регистрация делает организацию пробной ===")
tid, T = signup("trial@ex.io")
rec = main._tenant_rec(tid)
check(rec.get("trial") is True and rec.get("pagesCredit") == 1.0, "флаг trial и одна страница")

print("=== 2. Большой файл — фрагмент, а не 402 ===")
book = docx_of(40)                            # ≈1200 слов ≈ 4.8 стр.
r = up(T, book)
check(r.status_code == 200, "загрузка прошла: %s" % r.text[:120])
p = r.json()
ex = p.get("trialExcerpt") or {}
check(ex.get("filePages", 0) > 4 and 0 < ex.get("pages", 9) <= 1.0,
      "объём файла назван целиком, фрагмент — не больше страницы: %s" % ex)
check(len(p["segments"]) == ex["to"] - ex["from"] + 1 and ex["totalUnits"] == 40,
      "в проекте ровно строки фрагмента, подряд")
nums = [int(s["source"].split()[1]) for s in p["segments"]]
check(nums == list(range(nums[0], nums[0] + len(nums))), "строки фрагмента идут подряд, как в файле")
check(all(len(s["source"].split()) >= 8 for s in p["segments"]), "фрагмент непустой")
rec = main._tenant_rec(tid)
check(abs(rec["pagesUsed"] - ex["pages"]) < 1e-6 and rec["pagesUsed"] <= rec["pagesCredit"],
      "списан фрагмент, а не файл: %s" % rec["pagesUsed"])
log = rec.get("pagesLog") or []
check(any(e.get("kind") == "excerpt" and e.get("filePages", 0) > 4 for e in log),
      "строка журнала excerpt называет и объём файла")
pid = p["id"]
full = c.get("/api/projects/%d" % pid, headers=H(T)).json()
check(full.get("trialExcerpt") and not full.get("parseOutdated"), "отметка фрагмента доходит до браузера")

print("=== 3. Закрыто всё, что завело бы строки по всему файлу ===")
r = c.post("/api/projects/%d/resegment" % pid, headers=H(T), json={"dry_run": True})
check(r.status_code == 409 and "пробный фрагмент" in r.json()["detail"], "пересборка строк — 409")
r = c.post("/api/projects/%d/jobs" % pid, headers=H(T), json={"kind": "images", "segment_ids": []})
check(r.status_code == 409 and "пробный фрагмент" in r.json().get("detail", ""), "разбор картинок — 409: %s" % r.text[:100])
try:
    main._resegment_plan(main.get_project.__wrapped__(pid) if hasattr(main.get_project, "__wrapped__")
                         else next(x for x in main.STATE["projects"] if x["id"] == pid),
                         {"units": [], "full": []})
    check(False, "_resegment_plan (инструмент) пропустил фрагмент")
except main.HTTPException as e:
    check(e.status_code == 409, "и инструмент пересборки — 409")

print("=== 4. «Перевести остальное»: без пополнения — 402, после — весь файл ===")
r = c.post("/api/projects/%d/trial/rest" % pid, headers=H(T))
check(r.status_code == 402, "без пополнения — 402: %s" % r.text[:100])
check(len(next(x for x in main.STATE["projects"] if x["id"] == pid)["segments"]) == len(p["segments"]),
      "402 ничего не дописал")
r = c.post("/api/admin/tenants/%s" % tid, headers=H(A), json={"addPages": 10})
check(r.status_code == 200 and not main._tenant_rec(tid).get("trial"), "пополнение администратором снимает пробность")
r = c.post("/api/projects/%d/trial/rest" % pid, headers=H(T))
check(r.status_code == 200, "после пополнения — дописано: %s" % r.text[:120])
proj = next(x for x in main.STATE["projects"] if x["id"] == pid)
check(len(proj["segments"]) == 40 and not proj.get("trialExcerpt"), "в проекте весь файл, отметки фрагмента нет")
rec = main._tenant_rec(tid)
check(abs(rec["pagesUsed"] - ex["filePages"]) < 0.02, "всего списано ровно на файл: %s" % rec["pagesUsed"])
r = c.post("/api/projects/%d/resegment" % pid, headers=H(T), json={"dry_run": True})
check(r.status_code != 409 or "пробный" not in r.text, "пересборка снова открыта")

print("=== 5. Оплатившая организация фрагмента не получает ===")
r = up(T, docx_of(200, tag="глава"), "big.docx")     # ≈24 стр. при остатке ≈6
check(r.status_code == 402, "не пробная — прежний 402: %s" % r.text[:100])

print("=== 6. Маленький файл у пробной — обычная загрузка ===")
tid2, T2 = signup("small@ex.io")
r = up(T2, docx_of(3), "small.docx")
check(r.status_code == 200 and not r.json().get("trialExcerpt"), "файл меньше подарка — без фрагмента")

print("=== 7. Файл одним длинным абзацем — начало абзаца, а не 402 ===")
tid3, T3 = signup("long@ex.io")
d = Document()
d.add_paragraph(" ".join("Предложение номер %d про договор поставки товара." % i for i in range(120)))
b = io.BytesIO(); d.save(b)
r = up(T3, b.getvalue(), "contract.docx")
check(r.status_code == 200, "длинный абзац — загрузка прошла: %s" % r.text[:100])
ex3 = r.json().get("trialExcerpt") or {}
seg3 = r.json().get("segments") or []
check(ex3.get("cut") and len(seg3) == 1 and ex3.get("pages", 9) <= 1.0,
      "одна строка — начало абзаца, не больше страницы: %s" % {k: ex3.get(k) for k in ("pages", "filePages", "words")})
check(seg3 and seg3[0]["source"].endswith("."), "обрезано по предложению")
print()
print("ВСЁ ПРОШЛО" if not fail else "ПРОВАЛЕНО: " + "; ".join(fail))
sys.exit(1 if fail else 0)
