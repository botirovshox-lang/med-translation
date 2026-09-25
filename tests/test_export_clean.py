# -*- coding: utf-8 -*-
"""Выгрузка без внутренней кухни, тайминг в таблице субтитров, термины вне
прогона у субтитров, дорожка субтитров «показывать по умолчанию».

Что сторожится:
  1. xlsx: только №, оригинал (по галочке) и перевод — без статуса, маршрута,
     «риска»; строка с «=» остаётся текстом; в свойствах файла нет
     «openpyxl»; у субтитров — колонки начала и конца реплики;
  2. docx «просто текст»: нет числа сегментов, у субтитров — время в таблице,
     в свойствах нет «python-docx»;
  3. заголовки есть на каждом языке интерфейса (`UI_LANGS`);
  4. имя файла — код языка, а не «1в1» / «перевод»;
  5. у субтитров шаги терминов вне прогона по умолчанию — одно правило
     на разбор состава, `turnkey.params` и корзины; сегмент без termcheck
     не висит в «доделаю сама» и не зовётся «проверено начисто»;
     у документов всё как было;
  6. дорожка субтитров помечена `-disposition:s:0 default`.

Ни одного вызова модели.
"""
import io, os, sys, tempfile
from pathlib import Path
os.environ.setdefault("APP_PASSWORD", "test")
os.environ["AUTHORITY_CORPUS"] = "0"
os.environ.pop("OPENAI_API_KEY", None)
sys.path.insert(0, "backend")
import main
import media

main.save_state = lambda *a, **k: None
main.DEFAULT_UI_LANG = "ru"      # вне сессии заголовки — на языке по умолчанию
TMP = Path(tempfile.mkdtemp(prefix="medcat-expclean-"))
main.EXPORT_DIR = TMP / "exports"
fail = []


def check(cond, label):
    print(("  OK   " if cond else "  FAIL ") + label)
    if not cond:
        fail.append(label)


def proj(media_flag=False, segs=None):
    p = {"id": 1, "title": "Лекция", "src": "AR", "tgt": "UZ", "domain": "general",
         "tenant": main.DEFAULT_TENANT,
         "segments": segs or [
             {"id": 1, "source": "مرحبا", "target": "Salom", "status": "translated",
              "route": "GOOGLE_SAFE", "risk": "high", "provider": "gpt-5.6-sol"},
             {"id": 7, "source": "= 5", "target": "=HYPERLINK(\"x\")", "status": "qa",
              "route": "DUPLICATE", "risk": "low"}]}
    if media_flag:
        p["media"] = {"video": True}
        p["fileName"] = "lecture.srt"
    return p


# ─────────── 1. xlsx ───────────
print("=== 1. xlsx ===")
from openpyxl import load_workbook
real_cue_times = main._cue_times
main._cue_times = lambda p: None
path, _ = main._generate_export(proj(), "xlsx", include_source=True)
wb = load_workbook(str(path))
ws = wb.active
rows = [[c.value for c in r] for r in ws.iter_rows()]
flat = " ".join(str(v) for r in rows for v in r)
check(rows[0] == ["№", "Оригинал", "Перевод"], "шапка: №, оригинал, перевод — %r" % rows[0])
for bad in ("Статус", "Маршрут", "Риск", "GOOGLE_SAFE", "DUPLICATE", "high", "gpt-5.6", "qa"):
    check(bad not in flat, "в таблице нет «%s»" % bad)
check(rows[2][0] == 2, "номер — порядковый, а не id сегмента")
check(rows[2][2] == "=HYPERLINK(\"x\")" and ws.cell(3, 3).data_type == "s",
      "строка с «=» — текст, а не формула")
check((wb.properties.creator or "") == "" and "openpyxl" not in str(wb.properties.description or ""),
      "в свойствах файла нет openpyxl")
check(ws.title == "Лекция", "лист назван по проекту, а не «Segments»")
main.DEFAULT_UI_LANG = "uz"
check(main._export_rows(proj(), True)[0] == ["№", "Asl matn", "Tarjima"], "узбекский интерфейс — узбекская шапка")
main.DEFAULT_UI_LANG = "ru"

path, _ = main._generate_export(proj(), "xlsx", include_source=False)
ws = load_workbook(str(path)).active
check([c.value for c in ws[1]] == ["№", "Перевод"], "без галочки оригинала — только перевод")

main._cue_times = lambda p: {"1": [1.5, 3.25], "7": [3661.0, 3662.04]}
path, _ = main._generate_export(proj(media_flag=True), "xlsx", include_source=True)
rows = [[c.value for c in r] for r in load_workbook(str(path)).active.iter_rows()]
check(rows[0] == ["№", "Начало", "Конец", "Оригинал", "Перевод"], "у субтитров — тайминг: %r" % rows[0])
check(rows[1][1:3] == ["00:00:01,500", "00:00:03,250"], "время как в .srt: %r" % rows[1][1:3])
check(rows[2][1:3] == ["01:01:01,000", "01:01:02,040"], "часы считаются: %r" % rows[2][1:3])

# ─────────── 2. docx «просто текст» ───────────
print("\n=== 2. docx ===")
from docx import Document
path, _ = main._generate_export(proj(media_flag=True), "docx", include_source=False)
d = Document(str(path))
text = "\n".join(p.text for p in d.paragraphs)
check("сегмент" not in text and "экспорт" not in text, "в шапке нет числа сегментов и времени выгрузки")
check(d.tables and [c.text for c in d.tables[0].rows[0].cells] == ["№", "Начало", "Конец", "Перевод"],
      "у субтитров таблица с временем даже без оригинала")
check("python-docx" not in (d.core_properties.author or "") + (d.core_properties.comments or ""),
      "в свойствах нет python-docx")
main._cue_times = lambda p: None
path, _ = main._generate_export(proj(), "docx", include_source=False)
d = Document(str(path))
check(not d.tables and any(p.text == "Salom" for p in d.paragraphs),
      "документ без оригинала и тайминга — просто абзацы перевода")
main._cue_times = real_cue_times

# ─────────── 3. языки заголовков ───────────
print("\n=== 3. заголовки на каждом языке интерфейса ===")
check(set(main.UI_LANGS) <= set(main.EXPORT_HEADS), "EXPORT_HEADS покрывает UI_LANGS")
check(all(set(v) == set(main.EXPORT_HEADS["ru"]) for v in main.EXPORT_HEADS.values()),
      "у каждого языка те же ключи")

# ─────────── 4. имя файла ───────────
print("\n=== 4. имя файла ===")
p = proj(media_flag=True)
check(main._export_path(p, "docx_layout").name == "Лекция UZ.docx", main._export_path(p, "docx_layout").name)
check(main._export_path(p, "srt_bi").name == "Лекция AR-UZ.srt", main._export_path(p, "srt_bi").name)
check(main._export_path(p, "docx").name == "Лекция.docx", "обычный docx без хвоста")
odt = dict(proj(), importKind="odt", fileName="x.odt", writeback=False)
check(main._export_path(odt, "original") != main._export_path(odt, "docx_layout"),
      "«в исходном виде» Word-файлом не ложится в путь «как оригинал»")

# ─────────── 5. термины вне прогона у субтитров ───────────
print("\n=== 5. термины вне прогона у субтитров ===")
doc_p, sub_p = proj(), proj(media_flag=True)
check(main._default_run_steps(doc_p) == main.FULL_RUN_STEPS, "у документа — весь конвейер")
ds = main._default_run_steps(sub_p)
check("termcheck" not in ds and "termaudit" not in ds and "repair" in ds and ds[0] == "translate",
      "у субтитров без терминов, порядок прежний: %r" % ds)
check(main.FULL_RUN_STEPS[3:5] == ["termcheck", "termaudit"], "сам FULL_RUN_STEPS не тронут")

h = main._text_hash("Salom")
seg_ok = {"id": 1, "source": "مرحبا", "target": "Salom", "status": "translated",
          "backcheck": {"score": 99, "target_hash": h, "model": "m", "judged": True}}
check(main._machine_clean(seg_ok, 90) == main.CLEAN_NO_TERMCHECK, "донорство по-прежнему требует termcheck")
check(main._machine_clean(seg_ok, 90, need_tc=False) == main.CLEAN_TERMS_OFF,
      "для корзин субтитров — свой код, а не «начисто»")
seg_tc = dict(seg_ok, termcheck={"target_hash": h, "model": "m", "findings": [{"severity": "major"}]})
check(main._machine_clean(seg_tc, 90, need_tc=False) == main.CLEAN_TERMCHECK_FINDINGS,
      "свежие находки termcheck видны и у субтитров")

main.STATE = {"projects": [dict(sub_p, segments=[dict(seg_ok)])], "glossary": [], "tm": [],
              "termQueue": [], "exportHistory": [], "team": []}
main._ANALYSIS_CACHE.clear(); main._ANALYSIS_ROWS.clear()
main._invalidate_gloss_index()
main.get_project = lambda pid: main.STATE["projects"][0]
an = main.project_analysis(1, refresh=True)
tk = an["turnkey"]
check(1 not in tk["machine"], "сегмент без termcheck не висит в «доделаю сама» у субтитров")
check(1 in tk["ready"] and 1 in an["readyIds"], "он в «готово»")
check(1 not in an["clean"], "и не зовётся «проверено начисто»")
check(an["termsOff"] == {"inRun": False, "unchecked": 1}, "число для кнопки «Проверить термины»: %r" % an["termsOff"])
check(tk["params"]["steps"] == ds, "turnkey.params — то же правило")
rp = main.run_plan(1, main.RunPlanRequest())
check([p["step"] for p in rp["steps"]] == ds and rp["defaultSteps"] == ds,
      "разбор состава без шагов — то же умолчание, и оно названо")
rp = main.run_plan(1, main.RunPlanRequest(steps=["termcheck", "termaudit"]))
check([p["step"] for p in rp["steps"]] == ["termcheck", "termaudit"] and rp["defaultSteps"] == ds,
      "отдельная кнопка получает свои шаги")

main.STATE["projects"] = [dict(doc_p, segments=[dict(seg_ok)])]
main._ANALYSIS_CACHE.clear(); main._ANALYSIS_ROWS.clear()
an = main.project_analysis(1, refresh=True)
check(1 in an["turnkey"]["machine"], "у документа без termcheck — по-прежнему работа прогона")
check(an["termsOff"]["inRun"] is True and an["turnkey"]["params"]["steps"] == main.FULL_RUN_STEPS,
      "и полный конвейер")

# ─────────── 6. дорожка субтитров ───────────
print("\n=== 6. дорожка субтитров показывается сама ===")
seen = []
real_run = media.run
media.run = lambda cmd, **k: seen.append([str(x) for x in cmd])
media.mux_subtitles("in.mp4", "s.srt", "out.mp4", {"video": {"codec": "h264"}, "audio": {"codec": "aac"}}, "uzb")
media.mux_subtitles("in.mkv", "s.srt", "out.mkv", {"video": {"codec": "vp9"}}, "")
media.run = real_run
for cmd in seen:
    i = cmd.index("-disposition:s:0") if "-disposition:s:0" in cmd else -1
    check(i >= 0 and cmd[i + 1] == "default", "дорожка помечена default: " + cmd[-1])

print()
if fail:
    print("ПРОВАЛЕНО: %d" % len(fail))
    for f in fail:
        print("  - " + f)
    sys.exit(1)
print("ВСЁ ПРОШЛО")
