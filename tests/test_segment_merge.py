# -*- coding: utf-8 -*-
"""Границы строк руками: «Склеить со следующей», «Разъединить», «Разрезать».

Разбор файла иногда режет абзац на две строки (стык страниц, врезка) или
склеивает два в один; пересборка по правилам чинит не всё. Правила двери:
модель не зовётся; оригинал — склейка текстов как есть (перенос слова
снимается); подпись человека снимается со следом; счётчик перевода заново
наследуется; номер удалённой строки новой не достаётся; выгрузка 1в1
очищает абзац-хвост склейки и пишет в абзац разрезанной строки склейку
переводов частей — только когда переведены все; идущий прогон — 409;
у форматов, где строка — ячейка файла, склейки нет.
"""
import io, os, sys, json, tempfile
from pathlib import Path
os.environ.setdefault("APP_PASSWORD", "merge-pass-1")
os.environ["DATABASE_URL"] = ""
sys.path.insert(0, "backend")
import main
from docx import Document
from starlette.testclient import TestClient

main.save_state = lambda *a, **k: None
main._DICTIONARIES = []
main.STATE["users"], main.STATE["tenants"] = [], []
main.STATE["projects"] = []
main.STATE["glossary"] = []
main.STATE["folders"] = []
main.STATE["dicts"] = []
main._SESSIONS.clear(); main._LOGIN_FAILS.clear()
TMP = Path(tempfile.mkdtemp(prefix="medcat-merge-"))
main.SOURCE_DIR = TMP / "sources"
main.REIMPORT_DIR = TMP / "backups"
main.BOUNDARY_DIR = main.REIMPORT_DIR
main.EXPORT_DIR = TMP / "exports"
fail = []


def check(cond, label):
    print(("  OK   " if cond else "  FAIL ") + label)
    if not cond:
        fail.append(label)


c = TestClient(main.app)
H = lambda t: {"Authorization": "Bearer " + t}
main._ensure_users()
A = c.post("/api/auth/login", json={"login": "admin", "password": "merge-pass-1"}).json()["token"]

PARAS = ["Название этого продукта происходит от двух слов: «про» — впереди, «по-",
         "лис» — крепость, город. Это связано с летком улья.",
         "Прополис используется пчелами в качестве антисептического материала. Пчелы склеивают им соты.",
         "Свежесобранный прополис липкий и клейкий, со временем он твердеет.",
         "Последний абзац книги о пчелах и прополисе."]


def docx_bytes(paras):
    d = Document()
    for p in paras:
        d.add_paragraph(p)
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def live(pid):
    return next(p for p in main.STATE["projects"] if p["id"] == pid)


def export_texts(pid):
    out = TMP / ("out-%d.docx" % pid)
    stats = main._export_docx_layout(live(pid), out)
    return [p.text for p in Document(str(out)).paragraphs], stats


r = c.post("/api/projects/upload", headers=H(A), files={"file": ("book.docx", docx_bytes(PARAS))},
           data={"title": "book", "src": "RU", "tgt": "UZ", "domain": "general"})
check(r.status_code == 200, "файл загружен: %s %s" % (r.status_code, r.text[:120]))
pid = r.json()["id"]
P = live(pid)
check(len(P["segments"]) == 5, "пять строк: %d" % len(P["segments"]))
for k, s in enumerate(P["segments"]):
    s["target"] = ("po-" if k == 0 else "lis" if k == 1 else "t%d" % (k + 1))
    s["status"] = "translated"
    s["mtDone"] = True
P["segments"][0]["retranslations"] = 2
P["segments"][1].update({"status": "confirmed", "confirmedBy": 1, "confirmedAt": "2026-09-19 10:00"})
ids = [s["id"] for s in P["segments"]]

print("=== 1. Склеить со следующей ===")
r = c.post("/api/segments/%d/%d/merge-next" % (pid, ids[0]), headers=H(A))
check(r.status_code == 200, "склейка: %s %s" % (r.status_code, r.text[:160]))
segs = live(pid)["segments"]
s1 = segs[0]
check(len(segs) == 4 and ids[1] not in [s["id"] for s in segs], "вторая строка ушла")
check("«полис» — крепость" in s1["source"], "перенос слова снят: %r" % s1["source"][40:90])
check(s1["target"] == "polis", "перевод склеен тем же правилом переноса: %r" % s1["target"])
check(s1["status"] == "review", "склеенный перевод — на просмотр")
check(s1.get("retranslations") == 2 and s1.get("mtDone"), "счётчик перевода заново унаследован")
check((s1.get("unconfirmed") or {}).get("how") == "merge" and not s1.get("confirmedBy"),
      "подпись человека снята со следом")
mp = main._load_source_map(pid)
check(mp.get("tails") == [1], "абзац второй строки — хвост склейки: %s" % mp.get("tails"))
texts, st = export_texts(pid)
check(texts[0] == "polis" and texts[1] == "" and st.get("merged") == 1 and st.get("mismatch") == 0,
      "выгрузка: перевод в голове, хвост очищен, расхождений нет: %s %s" % (texts[:2], st.get("mismatch")))

print("=== 2. Разъединить ===")
live(pid)["segments"][0]["target"] = "правка"
r = c.post("/api/segments/%d/%d/unmerge" % (pid, ids[0]), headers=H(A))
check(r.status_code == 409, "правили после склейки — 409 без force: %s" % r.status_code)
r = c.post("/api/segments/%d/%d/unmerge?force=true" % (pid, ids[0]), headers=H(A))
check(r.status_code == 200, "разъединение с force: %s %s" % (r.status_code, r.text[:160]))
segs = live(pid)["segments"]
check([s["id"] for s in segs] == ids and segs[1].get("confirmedBy") == 1 and segs[0]["target"] == "po-",
      "строки как до склейки, подпись второй на месте")
mp = main._load_source_map(pid)
check(not mp.get("tails") and sorted(mp["pairs"]) == [[k, ids[k]] for k in range(5)], "карта как до склейки")
texts, st = export_texts(pid)
check(texts[:2] == ["po-", "lis"], "выгрузка — снова два абзаца: %s" % texts[:2])

print("=== 3. Разрезать с местом в переводе ===")
src = segs[2]["source"]
at = src.index("Пчелы склеивают")
segs[2]["target"] = "Birinchi gap. Ikkinchi gap."
r = c.post("/api/segments/%d/%d/split" % (pid, ids[2]), headers=H(A),
           json={"at": at, "target_at": len("Birinchi gap. ")})
check(r.status_code == 200, "разрезка: %s %s" % (r.status_code, r.text[:160]))
segs = live(pid)["segments"]
a, b = segs[2], segs[3]
check(a["source"].endswith("материала.") and b["source"].startswith("Пчелы склеивают"), "оригинал разрезан")
check(a["target"] == "Birinchi gap." and b["target"] == "Ikkinchi gap." and a["status"] == b["status"] == "review",
      "перевод разрезан там же, обе — на просмотр")
check(b["id"] > max(ids) and b.get("mtDone"), "номер новой строки — новый, счётчик унаследован")
texts, st = export_texts(pid)
check(texts[2] == "Birinchi gap. Ikkinchi gap.", "выгрузка: абзац — склейка переводов частей: %r" % texts[2])
b["target"] = ""
texts, st = export_texts(pid)
check(texts[2] == PARAS[2], "часть без перевода — абзац остаётся оригиналом целиком")
r = c.post("/api/segments/%d/%d/merge-next" % (pid, a["id"]), headers=H(A))
segs = live(pid)["segments"]
check(r.status_code == 200 and segs[2]["source"] == PARAS[2] and not main._load_source_map(pid).get("tails"),
      "склейка частей одного абзаца — снова одна строка, без хвоста")

print("=== 4. Разрезать без места в переводе ===")
seg = segs[3]
seg["target"] = "Yangi propolis."
r = c.post("/api/segments/%d/%d/split" % (pid, seg["id"]), headers=H(A),
           json={"at": seg["source"].index(",") + 1})
segs = live(pid)["segments"]
check(r.status_code == 200 and not segs[3]["target"] and not segs[4]["target"]
      and segs[3].get("prevTarget") == "Yangi propolis." and segs[4]["status"] == "new",
      "обе части пустые, прежний перевод — подсказкой: %s" % r.text[:120])
r = c.post("/api/segments/%d/%d/split" % (pid, seg["id"]), headers=H(A), json={"at": 0})
check(r.status_code == 400, "разрез с краю — 400")

print("=== 5. Номера не переиспользуются ===")
last = segs[-1]
r = c.post("/api/segments/%d/%d/merge-next" % (pid, segs[-2]["id"]), headers=H(A))
gone = last["id"]
seg = live(pid)["segments"][0]
r2 = c.post("/api/segments/%d/%d/split" % (pid, seg["id"]), headers=H(A), json={"at": 10})
new_ids = [s["id"] for s in live(pid)["segments"]]
check(r.status_code == 200 and r2.status_code == 200 and gone not in new_ids,
      "номер ушедшей последней строки новой не достался: %s" % gone)
r = c.post("/api/segments/%d/%d/merge-next" % (pid, live(pid)["segments"][-1]["id"]), headers=H(A))
check(r.status_code == 400, "у последней строки склеивать не с чем — 400")

print("=== 6. Запреты ===")
main._JOBS[999003] = {"id": 999003, "project": pid, "kind": "translate", "status": "running"}
r = c.post("/api/segments/%d/%d/merge-next" % (pid, live(pid)["segments"][0]["id"]), headers=H(A))
main._JOBS.pop(999003, None)
check(r.status_code == 409, "идущий прогон — 409: %s" % r.status_code)
live(pid)["slotsSha"] = "x"
r = c.post("/api/segments/%d/%d/merge-next" % (pid, live(pid)["segments"][0]["id"]), headers=H(A))
live(pid).pop("slotsSha")
check(r.status_code == 400, "формат со слотами — склейки нет: %s" % r.status_code)
img = {"id": 9001, "source": "Надпись", "target": "", "status": "new", "origin": {"kind": "image", "part": "p", "block": 0}}
live(pid)["segments"].append(img)
r = c.post("/api/segments/%d/9001/split" % pid, headers=H(A), json={"at": 3})
check(r.status_code == 400, "строку с картинки не режем: %s" % r.status_code)
c.post("/api/admin/tenants", headers=H(A),
       json={"id": "beta", "name": "Beta", "ownerLogin": "beta", "ownerPassword": "beta-pass-123"})
E = c.post("/api/auth/login", json={"login": "beta", "password": "beta-pass-123"}).json().get("token")
r = c.post("/api/segments/%d/%d/merge-next" % (pid, live(pid)["segments"][0]["id"]), headers=H(E))
check(E and r.status_code == 404, "чужая организация — 404: %s" % r.status_code)

print("=== 7. Пересборка называет ручные правки ===")
r = c.post("/api/projects/%d/resegment" % pid, headers=H(A), json={})
check(r.status_code == 200 and r.json().get("manualEdits", 0) > 0,
      "сухой прогон пересборки называет ручные правки границ: %s" % r.json().get("manualEdits"))

print("=== 8. Перевод половины и цепочка склеек ===")
PARAS2 = ["Первый абзац второй книги о меде.", "Второй абзац второй книги.",
          "Третий абзац второй книги о воске.", "Четвертый абзац второй книги.",
          "Пятый абзац второй книги о прополисе."]
r = c.post("/api/projects/upload", headers=H(A), files={"file": ("book2.docx", docx_bytes(PARAS2))},
           data={"title": "book2", "src": "RU", "tgt": "UZ", "domain": "general"})
pid2 = r.json()["id"]
segs = live(pid2)["segments"]
ids2 = [s["id"] for s in segs]
for k, s in enumerate(segs):
    s["target"] = "" if k == 1 else "U%d" % (k + 1)
    s["status"] = "new" if k == 1 else "translated"
r = c.post("/api/segments/%d/%d/merge-next" % (pid2, ids2[0]), headers=H(A))
s1 = live(pid2)["segments"][0]
check(r.status_code == 200 and s1["target"] == "" and s1["status"] == "new" and s1.get("prevTarget") == "U1",
      "переведена половина — строка новая, готовое подсказкой: %r %r" % (s1["target"], s1.get("prevTarget")))
texts, st = export_texts(pid2)
check(texts[0] == PARAS2[0] and texts[1] == PARAS2[1],
      "голова не записана — хвост НЕ очищен, оба абзаца оригиналом: %s" % texts[:2])
r = c.post("/api/segments/%d/%d/merge-next" % (pid2, ids2[2]), headers=H(A))
check(r.status_code == 200 and main._load_source_map(pid2).get("tails") == [1, 3], "вторая склейка — свой хвост")
r = c.post("/api/segments/%d/%d/unmerge" % (pid2, ids2[0]), headers=H(A))
check(r.status_code == 200 and main._load_source_map(pid2).get("tails") == [3],
      "откат первой склейки не трогает хвост второй: %s" % main._load_source_map(pid2).get("tails"))
texts, st = export_texts(pid2)
check(texts[2] == "U3 U4" and texts[3] == "", "вторая склейка выгружается как прежде: %s" % texts[2:4])

print("=== 9. Пересборка снимает ручные правки ===")
r = c.post("/api/projects/%d/resegment" % pid2, headers=H(A), json={"dry_run": False})
check(r.status_code == 200 and not any(s.get("boundary") for s in live(pid2)["segments"]),
      "после пересборки меток ручных правок нет: %s" % r.text[:120])

print()
if fail:
    print("FAILED: %d" % len(fail))
    for x in fail:
        print(" -", x)
    sys.exit(1)
print("ALL OK")
