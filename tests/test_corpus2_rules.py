"""Правила раздела Э (второй корпус выгрузки, файлы 41–70, 24.10.2026).

Разбор PDF проверяется на НАСТОЯЩЕЙ геометрии: файлы собираются рецептами
корпуса (`tools/corpus2_a.py`) во временный каталог и читаются тем же
путём, что и при загрузке (`textcount._pdf_pages_geom` → `pdftext.clean`).
Вёрстка и приёмка — на малых входах.

    python tests/test_corpus2_rules.py
"""
import os
import random
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))
sys.path.insert(0, str(ROOT / "backend"))     # backend первым: в tools есть свой export_qa.py
for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8")
    except Exception:
        pass

import pdftext  # noqa: E402
import textcount  # noqa: E402
import layout_pdf  # noqa: E402
import export_qa  # noqa: E402
import pdfpages_worker as W  # noqa: E402

FAIL = []


def check(ok, msg):
    print(("  OK   " if ok else "  FAIL ") + msg)
    if not ok:
        FAIL.append(msg)


def paras_of(rid: str) -> list:
    import corpus2_a
    fn = {r[0]: r[2] for r in corpus2_a.RECIPES}[rid]
    d = Path(tempfile.mkdtemp(prefix="c2r_"))
    p = d / (rid + ".pdf")
    random.seed(rid)
    fn(p)
    pages, geoms = textcount._pdf_pages_geom(p.read_bytes(), [])
    res = pdftext.clean(pages, geom=geoms)
    return [" ".join(str(it[1]).split()) for it in res["items"] if it and it[0] == "p"]


print("── Э1: подпись поля и значение — два места ──")
a = {"t": "Пациент:", "st": "b", "x0": 40.0, "x1": 80.0, "s": 8.5}
check(pdftext._field_pair(a, {"t": "Сорокина Е. В.", "st": "", "x0": 89.0, "s": 8.5}), "полужирная подпись с двоеточием и прямое значение")
check(not pdftext._field_pair(a, {"t": "Сорокина", "st": "b", "x0": 89.0, "s": 8.5}), "то же начертание — не делим")
check(not pdftext._field_pair(a, {"t": "текст продолжается", "st": "", "x0": 89.0, "s": 8.5}), "значение со строчной — продолжение фразы")
check(not pdftext._field_pair({"t": "Пациент", "st": "b", "x0": 40.0, "x1": 80.0}, {"t": "X", "st": "", "x0": 89.0}), "без двоеточия — не делим")
ps = paras_of("48_lab_results")
check("Пациент:" in ps and not any(p.startswith("Пациент: ") for p in ps),
      "бланк анализов: «Пациент:» отдельно от фамилии")
check(any(p.startswith("Сорокина") for p in ps), "значение поля — своё место")

print("── Э3–Э5: обтекание круглого фото и выключка словами ──")
ps = paras_of("43_magazine_wrap")
check(not any(p in ("в молочной", "Весной") for p in ps),
      "слова у края круглого фото не выпали из строки (Э3)")
check(any(p.startswith("Весной на плато") for p in ps), "абзац «Весной на плато…» цел")
check(any("высматривая добычу. Дети целыми днями" in p for p in ps),
      "«…добычу. Дети» — одна фраза, а не два абзаца (Э5)")
check(not any(p.startswith("косить приходится") for p in ps),
      "строка у края фото не стала красной строкой (Э4)")

print("── Э9–Э10: договор — пункты целиком ──")
ps = paras_of("44_contract")
check(not any(p.startswith("Заявка считается") for p in ps),
      "термин после конца фразы — не двойник надписи (Э9)")
check(not any(p.startswith("рублей за одну") for p in ps),
      "межсловный пробел соседних строк — свидетель (Э10)")
check(pdftext._twin_piece("На заметку", "На заметку $"), "двойник надписи по-прежнему узнаётся")
check(not pdftext._twin_piece("частью Договора (далее — «Заявка»).", "Заявка"), "повтор после точки — не двойник")

print("── Э2: ширина строки по профилю обтекания ──")
shape = [{"page": 0, "y": 100 - 12 * k, "x": 40.0, "x1": (190.0 if 3 <= k <= 5 else 288.0),
          "size": 9.0, "first": k == 0, "last": k == 9, "off": 0.0} for k in range(10)]
col = {"x0": 40.0, "x1": 288.0}
l_, r_ = layout_pdf._shape_at(shape, 0, 100 - 12 * 4, 12.0, col)
check(r_ == 190.0 and l_ is None, "строка у врезки получает правый предел врезки: %s" % ((l_, r_),))
l_, r_ = layout_pdf._shape_at(shape, 0, 100 - 12 * 8, 12.0, col)
check((l_, r_) == (None, None), "полная строка — без предела")
ragged = [dict(e, x1=288.0 - (k * 37 % 50)) for k, e in enumerate(shape)]
check(layout_pdf._shape_at(ragged, 0, 100 - 12 * 4, 12.0, col)[1] is None,
      "рваный край — короткая строка не фигура")
sl = [{"x": 40.0, "x1": 288.0, "y": 100 - 12 * k, "size": 9.0, "page": 0} for k in range(4)]
sl[-1] = {"x": 160.0, "x1": 230.0, "y": 64.0, "size": 9.0, "page": 0}
orig_vis = layout_pdf._vis_lines
layout_pdf._vis_lines = lambda s_: s_
try:
    check(layout_pdf._slot_align(sl) == "left", "Э2б: последнюю строку сдвинула фигура — не «по центру»")
finally:
    layout_pdf._vis_lines = orig_vis

print("── Э3: мягкое накрытие кривой обрезкой ──")
circle = ([[(100.0, 150.0), (150.0, 100.0), (100.0, 50.0), (50.0, 100.0)]], False)
ch = [("в", 55.0, 145.0, 9.0, False, "", "", False, 1),          # угол рамки круга
      ("молочной", 62.0, 145.0, 9.0, False, "", "", False, 1),
      ("пелене.", 160.0, 145.0, 9.0, False, "", "", False, 1)]  # за кругом — открыто
cov = [(5, [50.0, 50.0, 150.0, 150.0], (circle,))]
out = W._mark_covered(ch, cov)
check(not out[0][7] and not out[1][7], "слова в углу рамки круга рядом с открытым словом не накрыты")
ch2 = [("надпись", 95.0, 100.0, 9.0, False, "", "", False, 1)]
check(W._mark_covered(ch2, cov)[0][7], "слово внутри круга по-прежнему накрыто")

print("── Э7: повёрнутая строка — место по пути ──")
bx = {"x0": 40.0, "x1": 40.0, "top": 372.1, "bottom": 300.6, "size": 8.0,
      "slots": [[40.0, 302.6, 40.0, 8.0, "b", None, 11, 0,
                 {"pts": [[40.0, 302.6, 90.0, 8.0, 11, "ПОНЕДЕЛЬНИК"]], "end": [40.0, 365.3]}]]}
r = export_qa._rect_px(bx, 1000, 1400, 1.5, 842.0)
check(r is not None and r[2] - r[0] >= 20 and r[3] - r[1] >= 90, "рамка повёрнутой строки не пустая: %s" % (r,))

print("── Э8: номер пункта слева — якорь истории ──")
num = {"page": 0, "x": 40.0, "x1": 58.0, "y": 500.0, "size": 9.0}
fb = {"page": 0, "x": 70.0, "x1": 300.0, "y": 500.0, "size": 9.0}
B = {"runs": [[fb]]}
check(layout_pdf._num_anchor(fb, B, {0: [num, fb]}), "«3.3.» слева на той же строке — якорь")
far = dict(num, x=10.0, x1=20.0)
check(not layout_pdf._num_anchor(fb, B, {0: [far, fb]}), "далёкая надпись слева — не номер")

print("── Э12: первая строка абзаца и строка шапки — по геометрии ──")
import image_text  # noqa: E402
col = [{"box": [99, 550, 971, 582]}, {"box": [99, 595, 960, 627]}]
check(image_text._first_or_centered([148, 505, 1044, 539], col), "красная строка над колонкой спасена")
check(not image_text._first_or_centered([148, 505, 300, 539], col), "обрывок над колонкой — нет")
head = [{"box": [320, 85, 920, 125]}, {"box": [380, 175, 860, 200]}]
check(image_text._first_or_centered([363, 137, 879, 166], head), "строка шапки по центру между двумя спасена")

print("── Э14: наклонные строки — четырёхугольник, а не рамка ──")
check(abs(image_text._quad_angle([[0, 10], [100, 0], [100, 20], [0, 30]]) + 5.7) < 0.2, "угол по верхнему краю")
qa = {"box": [0, 0, 300, 80], "conf": 1.0, "flat": 0.0,
      "quad": [[0, 50], [300, 0], [304, 22], [4, 72]], "ang": -9.5}
qb = {"box": [0, 30, 300, 110], "conf": 1.0, "flat": 0.0,
      "quad": [[0, 80], [300, 30], [304, 52], [4, 102]], "ang": -9.5}
got = image_text._dedupe_boxes([qa, qb])
check(len(got) == 2, "соседние наклонные строки с перекрытыми рамками — две строки: %d" % len(got))
bl = image_text.group_blocks([dict(qa), dict(qb)])
check(len(bl) == 2 and all(b.get("quad") for b in bl), "наклонная строка — свой блок с четырёхугольником")

print("── Э13: печать и подпись — не переводятся ──")
os.environ.setdefault("APP_PASSWORD", "x")
os.environ.setdefault("WARM_CACHES", "0")
import io as _io, contextlib as _cl  # noqa: E402
with _cl.redirect_stdout(_io.StringIO()), _cl.redirect_stderr(_io.StringIO()):
    import main  # noqa: E402
sysp = main._image_read_system({"en": "general"}, "RU")
check("rubber stamp" in sysp and "signature" in sysp, "промпт чтения называет печать и подпись надпечаткой")

print("── Э15: сноски Word — в разборе, после колонтитулов ──")
import corpus2_b  # noqa: E402
from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
fb = {r[0]: r[2] for r in corpus2_b.RECIPES}
d_ = Path(tempfile.mkdtemp(prefix="c2r_"))
random.seed("58_docx_complex")
fb["58_docx_complex"](d_ / "58.docx")
doc = Document(str(d_ / "58.docx"))
allp = main._docx_flat_paragraphs(doc)
nonotes = main._docx_flat_parts(doc, notes=False)
txt = [main._docx_clean("".join(t.text or "" for t in p.iter(qn("w:t")))) if hasattr(main, "_docx_clean")
       else "".join(t.text or "" for t in p.iter(qn("w:t"))) for p in allp]
check(len(allp) > len(nonotes), "сноски добавлены к абзацам: %d > %d" % (len(allp), len(nonotes)))
check([p for p, _ in nonotes] == allp[:len(nonotes)], "номера тела и колонтитулов не сдвинулись")

print("── Э16: направление письма — по языку перевода ──")
random.seed("52_ar_letter")
fb["52_ar_letter"](d_ / "52.docx")
doc = Document(str(d_ / "52.docx"))
n_ = main._docx_flip_direction(doc, qn, "AR", "RU")
xml = doc.element.xml
check(n_ > 0 and "<w:bidi/>" not in xml and "w:bidiVisual" not in xml, "AR→RU: признаки справа налево сняты (%d)" % n_)
doc2 = Document(str(d_ / "52.docx"))
check(main._docx_flip_direction(doc2, qn, "AR", "FA") == 0, "оба языка справа налево — не трогаем")

print("── Э17: цифры арабского письма ──")
check(main._lang_normalize("Номер: ٤٨٧/٢٠٢٦", "RU") == "Номер: 487/2026", "RU — обычные цифры")
check(main._lang_normalize("رقم ٤٨٧", "AR") == "رقم ٤٨٧", "AR — как есть")

print("── Э18: диаграммы PowerPoint — в слотах ──")
import importers  # noqa: E402
random.seed("56_pt_slides")
fb["56_pt_slides"](d_ / "56.pptx")
pb = (d_ / "56.pptx").read_bytes()
s3 = importers.extract_slots("a.pptx", pb)["slots"]
s2 = importers.extract_slots("a.pptx", pb, 2)["slots"]
check(len(s3) > len(s2) and s3[:len(s2)] == s2, "тексты диаграммы — после слайдов, слайды на месте")
check(any("Toneladas" in t for t in s3[len(s2):]), "подпись оси диаграммы в слотах")
tr = {i: "RU " + t for i, t in enumerate(s3) if i >= len(s2) and t}
out = importers.write_back("a.pptx", pb, tr)
import zipfile  # noqa: E402
xmlc = zipfile.ZipFile(_io.BytesIO(out)).read("ppt/charts/chart1.xml").decode("utf-8")
check("RU Toneladas" in xmlc and "<c:v>RU Milho</c:v>" in xmlc, "перевод встал в диаграмму")

print("── Э20: лист снят в перспективе — блоки по выпрямленным строкам ──")
import math as _m
import numpy as _np
import image_text as IT  # noqa: E402
import text_paint as TP  # noqa: E402


def _ln(x0, yc, w, t, ang, bg=(200, 190, 170)):
    """Строка детектора: четырёхугольник ширины w и толщины t под углом ang
    от левого края на высоте yc, рамка — его ось-параллельный охват."""
    c, s_ = _m.cos(_m.radians(ang)), _m.sin(_m.radians(ang))
    q = [[x0, yc - t / 2], [x0 + w * c, yc - t / 2 + w * s_],
         [x0 + w * c, yc + t / 2 + w * s_], [x0, yc + t / 2]]
    xs, ys = [p_[0] for p_ in q], [p_[1] for p_ in q]
    return {"box": [int(min(xs)), int(min(ys)), int(max(xs)), int(max(ys))], "q4": q,
            "conf": 0.9, "flat": 0.9, "bg": list(bg)}


# абзац из пяти строк (толщина 28, шаг 36, наклон −3°) и под ним три ячейки
# таблицы строкой ниже — как «объявление в тени»
para = [_ln(200, 320 + 36 * i, 650, 28, -3.0) for i in range(5)]
cells = [_ln(190, 520, 110, 30, -3.0), _ln(370, 512, 130, 30, -2.5), _ln(625, 505, 70, 30, -2.4)]
check(IT._page_skew(para + cells) >= 1.0, "наклон листа узнан по длинным строкам")
bl = IT.group_blocks([dict(l) for l in para + cells])
rows = sorted(b["rows"] for b in bl)
check(rows == [1, 1, 1, 5], "абзац — один блок, ячейки таблицы — свои: %s" % rows)
big = [b for b in bl if b["rows"] == 5][0]
check(big["lineH"] <= 32, "высота строки — настоящая толщина, а не раздутая рамка: %s" % big["lineH"])
check(big["box"][1] <= min(l["box"][1] for l in para), "наружу — исходные рамки (по ним стирают)")
flat = [_ln(200, 320 + 36 * i, 650, 28, 0.0) for i in range(5)]
check(IT._page_skew(flat) == 0.0, "ровный лист — сборка как была")
# заголовок 41 и подзаголовок 28 на снимке под углом — разные блоки
hd = [_ln(390, 200, 340, 41, -3.4), _ln(305, 245, 510, 28, -3.4)] + para
check(len([b for b in IT.group_blocks([dict(l) for l in hd]) if b["rows"] == 1]) == 2,
      "заголовок и подзаголовок разного кегля — разные блоки")
# тень между строками абзаца — тот же фон (оттенок тот же, темнее)
sh = [dict(l, bg=[187, 177, 155]) for l in para[:2]] + [dict(l, bg=[140, 133, 114]) for l in para[2:]]
check(max(b["rows"] for b in IT.group_blocks(sh)) == 5, "граница тени абзац не режет")
check(not IT._same_light((240, 240, 240), (230, 160, 160)), "а розовая плашка — другой фон")
# наклонная линия таблицы между строками — разделитель
img = _np.full((120, 300, 3), 160, _np.uint8)
for x in range(300):
    y = int(round(60 + (x - 150) * _m.tan(_m.radians(-3.0))))
    img[y:y + 2, x] = 90
check(IT.divider_between(img, [40, 20, 260, 50], [40, 70, 260, 100], -3.0),
      "наклонная линия таблицы найдена вдоль наклона")
check(IT.divider_between(img, [40.4, 20.2, 259.7, 50.3], [40.1, 70.6, 260.2, 100.0], -3.0),
      "дробные (выпрямленные) рамки не роняют поиск")

print("── Э21: буквы в глубокой тени — порог краски от яркости бумаги ──")
from PIL import Image as _Im, ImageDraw as _Dr, ImageFont as _Ft
_fp = IT.font_path()
_fnt = _Ft.truetype(_fp, 30) if _fp else None
cv = _Im.new("RGB", (420, 60), (66, 64, 54))
_Dr.Draw(cv).text((10, 12), "Перед подачей воды", fill=(30, 27, 19), font=_fnt)
m = TP.measure(_np.array(cv), "", 1, 40.0)
check(m.get("found") and m.get("cover", 0) > 0.03, "тёмный текст на тёмной бумаге найден: %s" %
      {k: m.get(k) for k in ("found", "cover")})
cv2_ = _Im.new("RGB", (420, 60), (20, 40, 110))
_Dr.Draw(cv2_).text((10, 12), "ABC", fill=(40, 60, 130), font=_fnt)
m2 = TP.measure(_np.array(cv2_), "", 1, 40.0)
check(not m2.get("found"), "на цветной плашке порог прежний (почти невидимое — не «буквы»)")

# критик: строка из двух кусков на наклоне — высота настоящая, не раздутая
two = _ln(200, 320, 300, 28, -3.0)
two2 = _ln(520, 320 + 320 * _m.tan(_m.radians(-3.0)), 300, 28, -3.0)
j = dict(two, q4s=[two["q4"], two2["q4"]], box=[two["box"][0], two2["box"][1], two2["box"][2], two["box"][3]])
j.pop("q4")
r_ = IT._deskew_box(j)
check(r_[3] - r_[1] <= 32, "склеенная из кусков строка — высота одна толщина: %.1f" % (r_[3] - r_[1]))
check(not IT._same_light((200, 200, 200), (110, 110, 110)), "две серые разной яркости — не «тот же свет»")
bl2 = IT.group_blocks([dict(l) for l in para])
check(bl2[0]["lineBoxes"][0] == para[0]["box"] and "lineBoxesR" in bl2[0],
      "рамки строк — исходные, выпрямленные — отдельным полем для раскладки")
_rng = _np.random.default_rng(7)
night = _np.clip(_rng.normal(60, 9, (60, 420, 3)), 0, 255).astype(_np.uint8)
m3 = TP.measure(night, "", 1, 40.0)
check(not m3.get("found"), "шум ночного снимка на тёмном фоне — не буквы")

print("\n" + ("ВСЕ ПРОВЕРКИ ПРОШЛИ" if not FAIL else "ПРОВАЛЕНО: %d" % len(FAIL)))
sys.exit(1 if FAIL else 0)
