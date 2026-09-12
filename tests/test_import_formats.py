# -*- coding: utf-8 -*-
"""Импорт любого формата через .docx и повторный импорт того же файла.

Зачем. Конвейер проекта стоит на .docx (абзацы, якоря, исходник, экспорт
1в1, разбор картинок); чужой формат сначала превращается в .docx
(`backend/importers.py`), дальше идёт штатной дорогой. Повторный импорт:
человек правит документ и присылает снова — заводить новый файл значит
терять оплаченный перевод неизменившихся строк, а заменять по номерам —
сажать переводы на чужие строки. Поэтому сначала ПРОБА, потом диф по тексту
(`difflib`): совпавшее остаётся целиком, новое заводится на своём месте,
исчезнувшее уходит в копию для отката.

Что сторожится:
  1. txt/html/xlsx/png принимаются, строки разбираются, исходник .docx
     сохраняется рядом с оригиналом; картинка — 0 строк и пометка; html —
     по блочным тегам, <br> абзац не рвёт, строка таблицы — один абзац;
  2. неподдержанный формат — 415 с причиной; пустой — 415; те же байты
     под другим расширением разбираются заново (кэш — по sha И расширению);
  3. проба: тот же файл — exact; изменённый — similar с числами (в том числе
     заверённых и строк с картинок); короткий файл на книгу не похож;
     чужая папка — 404;
  4. повторный импорт: сухой прогон ничего не пишет; боевой сохраняет
     перевод и статус совпавших строк (и переехавших), обновляет оригинал
     при смене регистра, заводит новые на своём месте, убирает исчезнувшие
     и строки с картинок (в копию), списывает страницы только за добавленное;
     единица с двумя якорями даёт две строки карты на один сегмент; сегменты
     старого импорта (полный текст абзаца) узнаются по запасному ключу;
  5. откат возвращает прежние сегменты и исходник; после МАШИННОЙ правки
     (без отметок времени) — 409 без force; две замены подряд откатываются
     цепочкой; копия называется номером файла и уходит вместе с ним;
  6. удаление файла уносит оригинал и копии замен.

Ни одного вызова модели; файлы пишутся во временный каталог.
"""
import io, os, sys, shutil, tempfile
from pathlib import Path
os.environ["APP_PASSWORD"] = "import-pass-1"
os.environ["AUTHORITY_CORPUS"] = "0"
sys.path.insert(0, "backend")
import main
import importers
from starlette.testclient import TestClient

main.save_state = lambda *a, **k: None
main._DICTIONARIES = []
main.STATE["users"], main.STATE["tenants"] = [], []
main.STATE["projects"] = []
main.STATE["glossary"] = []
main.STATE["folders"] = []
main.STATE["dicts"] = []
main._SESSIONS.clear(); main._LOGIN_FAILS.clear()
TMP = Path(tempfile.mkdtemp(prefix="medcat-import-"))
main.SOURCE_DIR = TMP / "sources"
main.REIMPORT_DIR = TMP / "backups"
fail = []


def check(cond, label):
    print(("  OK   " if cond else "  FAIL ") + label)
    if not cond:
        fail.append(label)


c = TestClient(main.app)
H = lambda t: {"Authorization": "Bearer " + t}
main._ensure_users()
A = c.post("/api/auth/login", json={"login": "admin", "password": "import-pass-1"}).json()["token"]
c.post("/api/admin/tenants", headers=H(A),
       json={"id": "beta", "name": "Beta", "ownerLogin": "beta", "ownerPassword": "beta-pass-123"})
B = c.post("/api/auth/login", json={"login": "beta", "password": "beta-pass-123"}).json()["token"]


def upload(name, data, folder=None, tok=A):
    form = {"title": name.rsplit(".", 1)[0], "src": "RU", "tgt": "EN", "domain": "general"}
    if folder is not None:
        form["folder"] = str(folder)
    return c.post("/api/projects/upload", headers=H(tok), files={"file": (name, data)}, data=form)


def live(pid):
    return next(p for p in main.STATE["projects"] if p["id"] == pid)


print("=== 1. Форматы принимаются и разбираются ===")
TXT1 = "Первая строка.\nВторая строка.\nТретья строка.\nЧетвёртая строка.".encode("utf-8")
r = upload("doc.txt", TXT1)
check(r.status_code == 200, "txt принят: %s" % r.status_code)
P = r.json()
check(len(P["segments"]) == 4 and P["importKind"] == "txt" and P.get("sourceDocx"), "4 строки, kind=txt, исходник .docx записан")
check((main.SOURCE_DIR / ("%d.docx" % P["id"])).exists() and (main.SOURCE_DIR / ("%d.orig.txt" % P["id"])).exists(),
      "рядом лежат собранный .docx и оригинал .txt")
HTML = ("<html><body><h1>Заголовок</h1><p>Абзац <b>жирный</b> текст.</p><p>Улица Ленина, 5<br>Ташкент</p>"
        "<table><tr><td><p>А</p></td><td>Б</td></tr></table></body></html>").encode("utf-8")
r = upload("page.html", HTML)
check(r.status_code == 200 and [s["source"] for s in r.json()["segments"]]
      == ["Заголовок", "Абзац жирный текст.", "Улица Ленина, 5 Ташкент", "А | Б"],
      "html: блочные теги делят абзацы, инлайновые и <br> — нет, строка таблицы — один абзац")
from openpyxl import Workbook
wb = Workbook(); ws = wb.active; ws["A1"] = "Название товара"; ws["A2"] = "Хлеб ржаной"; ws["B2"] = 12
buf = io.BytesIO(); wb.save(buf)
r = upload("price.xlsx", buf.getvalue())
check(r.status_code == 200 and len(r.json()["segments"]) == 2 and "Word" in (r.json().get("importNote") or ""),
      "xlsx: текст ячеек стал строками, пометка про Word названа")
from PIL import Image
im = Image.new("RGB", (120, 80), "white"); b = io.BytesIO(); im.save(b, "PNG")
r = upload("scan.png", b.getvalue())
check(r.status_code == 200 and r.json()["segments"] == [] and r.json()["importKind"] == "image",
      "png: 0 строк, kind=image (текст читает разбор надписей)")
PNG_ID = r.json()["id"]
mp = main._load_source_map(PNG_ID)
check(mp is not None and mp.get("paras", 0) >= 1, "у картинки есть .docx-исходник с абзацем под картинку")

print("\n=== 2. Отказы формата и кэш разбора ===")
r = upload("virus.exe", b"MZ....")
check(r.status_code == 415 and "не поддерживается" in r.json()["detail"], "exe → 415 с причиной")
r = upload("empty.txt", b"")
check(r.status_code == 415, "пустой файл → 415: %s" % r.status_code)
SAME = "<p>один</p><p>два</p>".encode("utf-8")
p_txt = main._parse_upload("x.txt", SAME)
p_html = main._parse_upload("x.html", SAME)
check(p_txt["kind"] == "txt" and p_html["kind"] == "html" and len(p_html["paras"]) == 2 and len(p_txt["paras"]) == 1,
      "те же байты под .txt и .html разбираются по-разному (кэш по sha и расширению)")
check(main._parse_upload("x.html", SAME) is not None and main._PARSE_CACHE.get((p_html["sha"], ".html")) is not None
      and main._PARSE_CACHE[(p_html["sha"], ".html")][1]["docx"] is not None,
      "превращённый файл лежит в кэше вместе с .docx")
DOCX_P = main._parse_upload("plain.docx", importers.paragraphs_to_docx(["раз", "два"]))
check(main._PARSE_CACHE[(DOCX_P["sha"], ".docx")][1]["docx"] is None and DOCX_P["docx"] is not None,
      "у .docx собранный файл в кэше не хранится, но вызывающему отдаётся")

print("\n=== 3. Проба ===")
r = c.post("/api/projects/probe", headers=H(A), files={"file": ("doc.txt", TXT1)}, data={"folder": str(P["id"])})
check(r.status_code == 200 and [x["id"] for x in r.json()["exact"]] == [P["id"]], "тот же файл — exact")
TXT2 = "Первая строка.\nВторая строка исправлена.\nТретья строка.\nЧетвёртая строка.\nПятая новая.".encode("utf-8")
r = c.post("/api/projects/probe", headers=H(A), files={"file": ("doc.txt", TXT2)}, data={"folder": str(P["id"])})
s = r.json()["similar"]
check(r.json()["exact"] == [] and s and s[0]["id"] == P["id"] and s[0]["matched"] == 3 and s[0]["added"] == 2
      and s[0]["removed"] == 1 and "removedConfirmed" in s[0] and "images" in s[0],
      "изменённый файл — similar: совпало 3, новых 2, исчезнет 1, заверённые и картинки названы")
SHORT = "Первая строка.\nТретья строка.\nЧетвёртая строка.".encode("utf-8")
r = c.post("/api/projects/probe", headers=H(A), files={"file": ("short.txt", SHORT)}, data={"folder": str(P["id"])})
check(r.json()["similar"] and r.json()["similar"][0]["matched"] == 3, "три совпавшие строки из четырёх — ещё похоже")
TINY = "Первая строка.\nТретья строка.".encode("utf-8")
r = c.post("/api/projects/probe", headers=H(A), files={"file": ("tiny.txt", TINY)}, data={"folder": str(P["id"])})
check(r.json()["similar"] == [], "две строки на файл из четырёх — не похоже (порог и от старого, и от нового, и ≥ 3)")
r = c.post("/api/projects/probe", headers=H(B), files={"file": ("doc.txt", TXT2)}, data={"folder": str(P["id"])})
check(r.status_code == 404, "проба в чужую папку — 404")
r = c.post("/api/projects/probe", headers=H(B), files={"file": ("doc.txt", TXT1)})
check(r.status_code == 200 and r.json()["exact"] == [], "чужой файл не виден пробе другой организации")

print("\n=== 4. Повторный импорт ===")
proj = live(P["id"])
for sid, tgt, st in ((1, "First line.", "confirmed"), (3, "Third line.", "translated"), (4, "Fourth line.", "confirmed"), (2, "Second line.", "translated")):
    sg = next(x for x in proj["segments"] if x["id"] == sid)
    sg["target"], sg["status"] = tgt, st
    if st == "confirmed":
        sg["confirmedBy"] = 1
# сегмент с картинки — уходит в копию: его якорь описывает старый .docx
proj["segments"].append({"id": 9, "source": "Рис. 1", "target": "Fig. 1", "status": "translated",
                         "origin": {"kind": "image", "part": "word/media/image1.png", "block": 0}})
used_before = float((main._tenant_rec("default") or {}).get("pagesUsed") or 0)
r = c.post("/api/projects/%d/reimport" % P["id"], headers=H(A), files={"file": ("doc.txt", TXT2)}, data={"dry_run": "true"})
d = r.json()
check(r.status_code == 200 and d["dryRun"] and d["kept"] == 3 and d["added"] == 2 and d["removed"] == 1
      and d["images"] == 1 and len(proj["segments"]) == 5, "сухой прогон: числа названы, ничего не записано")
r = c.post("/api/projects/%d/reimport" % P["id"], headers=H(A), files={"file": ("empty2.txt", b"\n\n")}, data={"dry_run": "false"})
check(r.status_code == 415 and len(proj["segments"]) == 5, "файл без строк редакцией не считается — 415")
r = c.post("/api/projects/%d/reimport" % P["id"], headers=H(A), files={"file": ("doc.txt", TXT2)}, data={"dry_run": "false"})
d = r.json()
check(r.status_code == 200 and d["stamp"].startswith("%d-" % P["id"]) and d["kept"] == 3 and d["added"] == 2
      and d["removed"] == 1 and d["imagesRemoved"] == 1,
      "замена: " + str({k: d.get(k) for k in ("kept", "added", "removed", "imagesRemoved", "stamp")}))
segs = proj["segments"]
by_src = {s["source"]: s for s in segs}
check([s["source"] for s in segs] == ["Первая строка.", "Вторая строка исправлена.", "Третья строка.", "Четвёртая строка.", "Пятая новая."],
      "порядок — как в новом файле, строк с картинок нет")
check(by_src["Первая строка."]["id"] == 1 and by_src["Первая строка."]["target"] == "First line." and by_src["Первая строка."]["status"] == "confirmed",
      "совпавшая строка сохранила номер, перевод и заверение")
check(by_src["Вторая строка исправлена."]["status"] == "new" and by_src["Вторая строка исправлена."]["target"] == ""
      and by_src["Вторая строка исправлена."]["id"] > 9, "изменённая строка — новая, без перевода, свежий номер")
check(proj.get("reimport", {}).get("stamp") == d["stamp"] and proj["sourceSha"] != P["sourceSha"]
      and proj["reimport"]["addedIds"] == d["addedIds"], "отметка замены с номерами добавленных и новый sha")
mp = main._load_source_map(P["id"])
check(mp and len(mp["pairs"]) == 5 and "images" not in mp, "карта исходника переписана, карта картинок старого файла снята")
used_after = float((main._tenant_rec("default") or {}).get("pagesUsed") or 0)
check(used_after - used_before <= 1.0 + 1e-9 and d["pagesDebited"] < 1.0, "списано только за добавленные строки: %.3f" % d["pagesDebited"])
log = (main._tenant_rec("default") or {}).get("pagesLog") or []
check(log and log[-1]["kind"] == "reimport", "в журнале страниц — запись reimport")
check((main.REIMPORT_DIR / ("reimport-%s.json" % d["stamp"])).exists() and (main.REIMPORT_DIR / ("reimport-%s.docx" % d["stamp"])).exists()
      and list(main.REIMPORT_DIR.glob("reimport-%s.orig.*" % d["stamp"])), "копия для отката: сегменты, прежний исходник и оригинал")
STAMP1 = d["stamp"]

print("\n=== 4b. Переезд, повтор с двумя якорями, регистр, старый импорт ===")
by_src["Пятая новая."]["target"] = "Fifth new."
TXT3 = "Пятая новая.\nПЕРВАЯ СТРОКА.\nВторая строка исправлена.\nТретья строка.\nТретья строка.\nЧетвёртая строка.".encode("utf-8")
r = c.post("/api/projects/%d/reimport" % P["id"], headers=H(A), files={"file": ("doc.txt", TXT3)}, data={"dry_run": "false"})
d2 = r.json()
segs = proj["segments"]
check(r.status_code == 200 and d2["moved"] == 1 and d2["added"] == 0 and d2["removed"] == 0
      and segs[0]["source"] == "Пятая новая." and segs[0]["target"] == "Fifth new.",
      "переехавший абзац сохранил перевод: " + str({k: d2.get(k) for k in ("kept", "moved", "added", "removed")}))
check(segs[1]["id"] == 1 and segs[1]["source"] == "ПЕРВАЯ СТРОКА." and segs[1]["target"] == "First line.",
      "смена регистра: сегмент тот же, оригинал обновлён, перевод остался")
mp = main._load_source_map(P["id"])
third = [pr for pr in mp["pairs"] if pr[1] == 3]
check(len(third) == 2 and len(segs) == 5, "соседний повтор — один сегмент с двумя якорями в карте")
old_style = {"id": 777, "tenant": "default", "src": "RU", "tgt": "EN", "domain": "general",
             "segments": [{"id": 1, "source": "Введение 5", "target": "Introduction", "status": "translated"},
                          {"id": 2, "source": "Глава 1", "target": "Chapter 1", "status": "translated"}]}
plan, removed = main._diff_units(old_style, [("Введение", [0]), ("Глава 1", [1])], full=["Введение 5", "Глава 1"])
check([p[0] for p in plan] == ["keep", "keep"] and removed == [], "сегмент старого импорта (полный текст абзаца) узнаётся по запасному ключу")

print("\n=== 5. Откат ===")
STAMP2 = d2["stamp"]
r = c.post("/api/projects/%d/reimport/%s/undo" % (P["id"], STAMP1), headers=H(A))
check(r.status_code == 409, "откат не последней замены — 409")
segs[2]["target"] = "Machine wrote this."          # машинная правка: без отметок времени
r = c.post("/api/projects/%d/reimport/%s/undo" % (P["id"], STAMP2), headers=H(A))
check(r.status_code == 409, "после машинной правки откат без force — 409")
r = c.post("/api/projects/%d/reimport/%s/undo?force=true" % (P["id"], STAMP2), headers=H(A))
check(r.status_code == 200 and [s["source"] for s in proj["segments"]][:2] == ["Первая строка.", "Вторая строка исправлена."]
      and proj.get("reimport", {}).get("stamp") == STAMP1, "force: вернулось состояние после ПЕРВОЙ замены, отметка — её")
# Между первой и второй заменой у добавленной строки появился перевод
# («Fifth new.» — машинная правка без отметки времени): откат первой
# замены его унёс бы, значит 409, и только с force.
r = c.post("/api/projects/%d/reimport/%s/undo" % (P["id"], STAMP1), headers=H(A))
check(r.status_code == 409, "перевод у добавленной строки — тоже «менялось после замены»: 409")
r = c.post("/api/projects/%d/reimport/%s/undo?force=true" % (P["id"], STAMP1), headers=H(A))
check(r.status_code == 200 and [s["source"] for s in proj["segments"]][:4] == ["Первая строка.", "Вторая строка.", "Третья строка.", "Четвёртая строка."]
      and any((s.get("origin") or {}).get("kind") == "image" for s in proj["segments"]),
      "второй откат с force: прежние строки и сегмент картинки вернулись")
check(proj["sourceSha"] == P["sourceSha"] and "reimport" not in proj, "sha и отметка замены вернулись к прежним")
mp = main._load_source_map(P["id"])
check(mp and len(mp["pairs"]) == 4 and (main.SOURCE_DIR / ("%d.orig.txt" % P["id"])).exists(), "карта исходника и оригинал вернулись прежними")
r = c.post("/api/projects/%d/reimport/%s/undo" % (P["id"], STAMP1), headers=H(B))
check(r.status_code == 404, "откат чужому — 404")

print("\n=== 6. Удаление уносит оригинал и копии замен ===")
r = c.delete("/api/projects/%d" % P["id"], headers=H(A))
check(r.status_code == 200 and not list(main.SOURCE_DIR.glob("%d.*" % P["id"]))
      and not list(main.REIMPORT_DIR.glob("reimport-%d-*" % P["id"])), "файлы исходника, оригинала и копий удалены")

shutil.rmtree(TMP, ignore_errors=True)
print()
print("FAILED: %d" % len(fail) if fail else "ВСЁ ПРОШЛО")
for f_ in fail:
    print("  - " + f_)
sys.exit(1 if fail else 0)
